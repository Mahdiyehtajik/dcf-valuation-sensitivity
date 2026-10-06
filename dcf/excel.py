"""
Write the model to an Excel workbook with LIVE formulas, not pasted values.

A reviewer should be able to change an assumption in the Inputs sheet and see
the valuation move. The Python model is the source of truth; the workbook
mirrors its logic cell by cell, and `verify_against_python()` checks that the
two agree before the file is delivered.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .model import Assumptions, value

BLUE_INPUT = Font(color="0000FF")          # convention: blue = hard-coded input
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

INPUT_ROWS = [
    ("Last historical revenue", "revenue_0", "#,##0.0"),
    ("Revenue growth", "revenue_growth", "0.0%"),
    ("EBIT margin", "ebit_margin", "0.0%"),
    ("Tax rate", "tax_rate", "0.0%"),
    ("D&A % revenue", "da_pct", "0.0%"),
    ("Capex % revenue", "capex_pct", "0.0%"),
    ("NWC % revenue", "nwc_pct", "0.0%"),
    ("Risk-free rate", "risk_free", "0.00%"),
    ("Equity risk premium", "equity_risk_premium", "0.00%"),
    ("Beta", "beta", "0.00"),
    ("Pre-tax cost of debt", "cost_of_debt", "0.00%"),
    ("Debt weight D/(D+E)", "debt_weight", "0.0%"),
    ("Terminal growth", "terminal_growth", "0.00%"),
    ("Net debt", "net_debt", "#,##0.0"),
    ("Shares outstanding", "shares_outstanding", "#,##0.0"),
]


def write_workbook(a: Assumptions, path: str | Path, two_way=None) -> Path:
    wb = Workbook()
    inp = wb.active
    inp.title = "Inputs"
    inp["A1"], inp["B1"] = "Assumption", "Value"
    inp["A1"].font = inp["B1"].font = BOLD
    cell_of: dict[str, str] = {}
    for i, (label, field, fmt) in enumerate(INPUT_ROWS, start=2):
        inp.cell(i, 1, label)
        v = getattr(a, field)
        c = inp.cell(i, 2, v if v is not None else 0)
        c.font, c.number_format, c.border = BLUE_INPUT, fmt, BOX
        cell_of[field] = f"Inputs!$B${i}"
    r = len(INPUT_ROWS) + 3
    inp.cell(r, 1, "Cost of equity (CAPM)").font = BOLD
    inp.cell(r, 2, f"={cell_of['risk_free']}+{cell_of['beta']}*{cell_of['equity_risk_premium']}").number_format = "0.00%"
    cell_of["cost_of_equity"] = f"Inputs!$B${r}"
    inp.cell(r + 1, 1, "WACC").font = BOLD
    inp.cell(r + 1, 2, f"=(1-{cell_of['debt_weight']})*{cell_of['cost_of_equity']}+{cell_of['debt_weight']}*{cell_of['cost_of_debt']}*(1-{cell_of['tax_rate']})").number_format = "0.00%"
    cell_of["wacc"] = f"Inputs!$B${r + 1}"
    inp.column_dimensions["A"].width = 28
    inp.column_dimensions["B"].width = 14

    # ---- Model sheet: years across columns
    m = wb.create_sheet("Model")
    n = a.years
    cols = [get_column_letter(c) for c in range(2, 2 + n + 1)]  # B = year 0
    m["A1"] = "Year"
    m["A1"].font = BOLD
    for j, col in enumerate(cols):
        m[f"{col}1"] = j
        m[f"{col}1"].font = BOLD
        m[f"{col}1"].fill = HEAD_FILL
    lines = ["Revenue", "EBIT", "NOPAT", "D&A", "Capex", "NWC", "Change in NWC", "FCFF", "Discount factor", "PV of FCFF"]
    for i, name in enumerate(lines, start=2):
        m.cell(i, 1, name)
    rev, ebit, nopat, da, capex, nwc, dnwc, fcff, dfac, pv = range(2, 12)
    m[f"B{rev}"] = f"={cell_of['revenue_0']}"
    m[f"B{nwc}"] = f"=B{rev}*{cell_of['nwc_pct']}"
    for j in range(1, n + 1):
        col, prev = cols[j], cols[j - 1]
        m[f"{col}{rev}"] = f"={prev}{rev}*(1+{cell_of['revenue_growth']})"
        m[f"{col}{ebit}"] = f"={col}{rev}*{cell_of['ebit_margin']}"
        m[f"{col}{nopat}"] = f"={col}{ebit}*(1-{cell_of['tax_rate']})"
        m[f"{col}{da}"] = f"={col}{rev}*{cell_of['da_pct']}"
        m[f"{col}{capex}"] = f"={col}{rev}*{cell_of['capex_pct']}"
        m[f"{col}{nwc}"] = f"={col}{rev}*{cell_of['nwc_pct']}"
        m[f"{col}{dnwc}"] = f"={col}{nwc}-{prev}{nwc}"
        m[f"{col}{fcff}"] = f"={col}{nopat}+{col}{da}-{col}{capex}-{col}{dnwc}"
        t = f"({col}1-0.5)" if a.mid_year else f"{col}1"
        m[f"{col}{dfac}"] = f"=(1+{cell_of['wacc']})^(-{t})"
        m[f"{col}{pv}"] = f"={col}{fcff}*{col}{dfac}"
    for row in range(2, 12):
        for col in ["B"] + cols:
            m[f"{col}{row}"].number_format = "0.0000" if row == dfac else "#,##0.0"
    last = cols[-1]

    # ---- Valuation bridge
    v0 = 14
    bridge = [
        ("Sum of PV (explicit period)", f"=SUM(C{pv}:{last}{pv})"),
        ("Terminal value (Gordon)", f"={last}{fcff}*(1+{cell_of['terminal_growth']})/({cell_of['wacc']}-{cell_of['terminal_growth']})"),
        ("PV of terminal value", f"=B{v0 + 1}*(1+{cell_of['wacc']})^(-{last}1)"),
        ("Enterprise value", f"=B{v0}+B{v0 + 2}"),
        ("Less: net debt", f"={cell_of['net_debt']}"),
        ("Equity value", f"=B{v0 + 3}-B{v0 + 4}"),
        ("Value per share", f"=IF({cell_of['shares_outstanding']}>0,B{v0 + 5}/{cell_of['shares_outstanding']},\"n/a\")"),
        ("TV as % of EV", f"=B{v0 + 2}/B{v0 + 3}"),
    ]
    for i, (label, f) in enumerate(bridge):
        m.cell(v0 + i, 1, label).font = BOLD
        c = m.cell(v0 + i, 2, f)
        c.number_format = "0.0%" if "%" in label else "#,##0.0"
        c.border = BOX
    m.column_dimensions["A"].width = 28

    # ---- Optional sensitivity sheet (values from Python, since Excel data
    # tables cannot be created by openpyxl)
    if two_way is not None:
        s = wb.create_sheet("Sensitivity")
        s["A1"] = f"Enterprise value: {two_way.index.name} (rows) x {two_way.columns.name} (columns)"
        s["A1"].font = BOLD
        s.cell(3, 1, two_way.index.name).font = BOLD
        for j, cv in enumerate(two_way.columns, start=2):
            c = s.cell(3, j, cv)
            c.number_format, c.font, c.fill = "0.0%", BOLD, HEAD_FILL
        for i, rv in enumerate(two_way.index, start=4):
            c = s.cell(i, 1, rv)
            c.number_format, c.font, c.fill = "0.0%", BOLD, HEAD_FILL
            for j, cv in enumerate(two_way.columns, start=2):
                x = two_way.loc[rv, cv]
                s.cell(i, j, None if x != x else float(x)).number_format = "#,##0"
        s["A2"] = "Static values computed by the Python model (Excel data tables are not written by openpyxl)."
        s["A2"].font = Font(italic=True, color="808080")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def verify_against_python(a: Assumptions, path: str | Path, tol: float = 1e-6) -> float:
    """
    Recalculate the workbook with LibreOffice (headless) and compare the
    enterprise value cell to the Python model. Returns the absolute difference.
    Raises if LibreOffice is not available.
    """
    import subprocess, tempfile, shutil
    from openpyxl import load_workbook

    path = Path(path)
    tmp = Path(tempfile.mkdtemp())
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp), str(path)],
        check=True, capture_output=True,
    )
    wb = load_workbook(tmp / path.name, data_only=True)
    ev_xl = wb["Model"]["B17"].value
    ev_py = value(a)["enterprise_value"]
    shutil.rmtree(tmp, ignore_errors=True)
    diff = abs(ev_xl - ev_py)
    if diff > tol * max(1.0, abs(ev_py)):
        raise AssertionError(f"Excel EV {ev_xl} != Python EV {ev_py}")
    return diff
